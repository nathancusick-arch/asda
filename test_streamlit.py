"""Asda weekly report generator.

Upload an old/new audit CSV and the previous client report (.xlsx).
The LIVE report is not required. Reports contain calculated values only.
"""

from copy import copy
from datetime import date, datetime, time, timedelta
from decimal import Decimal, ROUND_HALF_UP
import io
import re
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl import load_workbook


SHEET_NAME = "Weekly Data Capture"
REPORT_HEADERS = [
    "Visit Ref", "Visit Type", "Product Category", "Date of Visit", "Time of Visit",
    "Location Number", "Store Name", "Postcode", "Asked for ID",
    "Product purchased without providing ID", "Type Of Till", "Queue at the till",
    "People in the queue", "Colleague made eye contact", "Supervisor called",
    "Colleague Gender", "Colleague Age", "Colleague Name", "Colleague Description",
    "Description of Item", "ST code", "OP code", "TE code", "TR code",
    "Visitor Comments", "SL Ref", "Did the staff member approach you to attend to your till?",
]
SUPPORTED_ITEMS = {
    "Alcohol - No ID", "Cigarettes - No ID", "Click & Collect", "E-Cig - No ID",
    "Knife - No ID", "Lottery - No ID", "Fireworks - No ID",
    "Supermarket Home Delivery", "Home Delivery",
}
VISIT_TYPES = {
    "scheduled_off_monthly": "Quarterly", "scheduled_monthly": "Quarterly",
    "revisit": "Revisits", "weekly": "Revisits", "emergency": "Emergency",
    "emergency_weekly": "Emergency",
}
DELIVERY_ITEMS = {"Supermarket Home Delivery", "Home Delivery", "Click & Collect"}
CHECKOUT_ID = "Did the store colleague who served you at the checkout ask you for ID?"
CABINET_ID = "If you were able to complete your purchase of the restricted item, did the colleague who served you at the fireworks cabinet, ask you for ID?"
ALLOW = "Did the store colleague allow you to purchase the restricted item without providing ID?"
HANDOVER = "Did the store colleague who served you at the fireworks cabinet, hand over the restricted item without providing ID?"

# The original mapper's question aliases, with surrounding whitespace normalised.
QUESTION_MAPPING = {'What type of alcohol did you try to purchase?': ['What type of alcohol did you purchase/try to purchase?',
                                                   'What type of E-cigarette product did you purchase/attempt to '
                                                   'purchase?',
                                                   'What type of alcohol did you try to purchase?'],
 'Please give details of the product that you purchased:': ['Please give details of the alcohol purchased:',
                                                            'Please give details of the lottery ticket you tried to '
                                                            'purchase/purchased:',
                                                            'Please give details of the cigarettes that you tried to '
                                                            'purchase/purchased:',
                                                            'Please give details of the alcohol that you tried to '
                                                            'purchase:',
                                                            'Please give details of the knife product that you tried '
                                                            'to purchase/purchased:',
                                                            'Please give details of the e-cig product that you tried '
                                                            'to purchase:',
                                                            'Please give details of the alcohol that you purchased:',
                                                            'Please give details of the fireworks that you tried to '
                                                            'purchase:'],
 'At which type of till was the purchase made?': 'At which type of till was the purchase made?',
 'Did the staff member approach you to attend to your till?': 'Did the staff member approach you to attend to your '
                                                              'till?',
 'Did you make the purchase on your own?': 'Did you make the purchase on your own?',
 'Did the store colleague who served you ask your age?': 'Did the store colleague who served you ask your age?',
 'Did the store colleague who served you ask you for ID?': ['Did the driver ask for ID?',
                                                            'Did the store colleague who served you ask you for ID?',
                                                            'Did the store colleague who served you ask for your ID?',
                                                            'Did the store colleague who served you ask for and '
                                                            'attempt to check your ID?',
                                                            'merged_id_asked'],
 'Please confirm that you did not present any ID:': 'Please confirm that you did not present any ID:',
 'Did the store colleague allow you to purchase the restricted item without providing ID?': ['Did the store colleague '
                                                                                             'allow you to purchase '
                                                                                             'the restricted item '
                                                                                             'without providing ID?',
                                                                                             'merged_no_id_allow_handover'],
 'Did the store colleague who served you make eye contact with you during the transaction?': ['Did the driver make eye '
                                                                                              'contact with you during '
                                                                                              'the interaction?',
                                                                                              'Did the store colleague '
                                                                                              'who served you make eye '
                                                                                              'contact with you during '
                                                                                              'the transaction?',
                                                                                              'Did the person who '
                                                                                              'served you make eye '
                                                                                              'contact with you during '
                                                                                              'the collection of your '
                                                                                              'shopping?'],
 'When was eye contact first made?': 'When was eye contact first made?',
 'What was the gender of the store colleague who served you?': ['What was the gender of the driver?',
                                                                'What was the gender of the store colleague who served '
                                                                'you?'],
 'What was the approximate age of the store colleague who served you?': 'What was the approximate age of the store '
                                                                        'colleague who served you?',
 'Please input the approximate height of the server:': ['Please accurately describe the driver:',
                                                        'Please input the approximate height of the server:',
                                                        'Please accurately describe the store colleague who served '
                                                        'you:'],
 "Please input the length of the server's hair:": "Please input the length of the server's hair:",
 'Please input the style of the server hair:': 'Please input the style of the server hair:',
 'Please input the hair colour of the server:': 'Please input the hair colour of the server:',
 'Please select whether the server had any of the following features:': 'Please select whether the server had any of '
                                                                        'the following features:',
 'Please input their distinguishing characteristics:': 'Please input their distinguishing characteristics:',
 'What was the name of the store colleague who served you?': 'What was the name of the store colleague who served you?',
 'How many people were in the queue?': 'How many people were in the queue?',
 "Were 'Challenge 25' notices visible at the till?": "Were 'Challenge 25' notices visible at the till?",
 'From the receipt, please enter the 4-digit ST code:': ['On the delivery receipt there is a list of order details, '
                                                         "please record the 'Order No':",
                                                         'From the receipt, please enter the 5-digit ST code:',
                                                         "Enter the 'Order No.':"],
 'Please enter the 8-digit OP code:': 'Please enter the 6-digit OP code:',
 'Please enter the 2-digit TE code:': 'Please enter the 2-digit TE code:',
 'Please enter the 5-digit TR code:': 'Please enter the 5-digit TR code:',
 'Please enter the time from the receipt:': 'Please enter the time from the receipt:',
 'Please use this space to explain anything unusual about your visit or to clarify any detail of your report:': 'Please '
                                                                                                                'use '
                                                                                                                'this '
                                                                                                                'space '
                                                                                                                'to '
                                                                                                                'explain '
                                                                                                                'anything '
                                                                                                                'unusual '
                                                                                                                'about '
                                                                                                                'your '
                                                                                                                'visit '
                                                                                                                'or to '
                                                                                                                'clarify '
                                                                                                                'any '
                                                                                                                'detail '
                                                                                                                'of '
                                                                                                                'your '
                                                                                                                'report:',
 'Please confirm in the space below whether or not you were asked for ID, and if so, whether the store colleague who served you allowed the transaction through without you presenting ID:': ['Please '
                                                                                                                                                                                              'confirm '
                                                                                                                                                                                              'below '
                                                                                                                                                                                              'whether '
                                                                                                                                                                                              'or '
                                                                                                                                                                                              'not '
                                                                                                                                                                                              'you '
                                                                                                                                                                                              'were '
                                                                                                                                                                                              'asked '
                                                                                                                                                                                              'for '
                                                                                                                                                                                              'ID:',
                                                                                                                                                                                              'Please '
                                                                                                                                                                                              'confirm '
                                                                                                                                                                                              'whether '
                                                                                                                                                                                              'or '
                                                                                                                                                                                              'not '
                                                                                                                                                                                              'you '
                                                                                                                                                                                              'were '
                                                                                                                                                                                              'asked '
                                                                                                                                                                                              'for '
                                                                                                                                                                                              'ID, '
                                                                                                                                                                                              'and '
                                                                                                                                                                                              'if '
                                                                                                                                                                                              'so, '
                                                                                                                                                                                              'at '
                                                                                                                                                                                              'what '
                                                                                                                                                                                              'point '
                                                                                                                                                                                              'during '
                                                                                                                                                                                              'the '
                                                                                                                                                                                              'transaction '
                                                                                                                                                                                              'ID '
                                                                                                                                                                                              'was '
                                                                                                                                                                                              'requested:',
                                                                                                                                                                                              'Please '
                                                                                                                                                                                              'confirm '
                                                                                                                                                                                              'whether '
                                                                                                                                                                                              'or '
                                                                                                                                                                                              'not '
                                                                                                                                                                                              'your '
                                                                                                                                                                                              'ID '
                                                                                                                                                                                              'was '
                                                                                                                                                                                              'checked, '
                                                                                                                                                                                              'and '
                                                                                                                                                                                              'if '
                                                                                                                                                                                              'so, '
                                                                                                                                                                                              'at '
                                                                                                                                                                                              'what '
                                                                                                                                                                                              'point '
                                                                                                                                                                                              'during '
                                                                                                                                                                                              'the '
                                                                                                                                                                                              'transaction '
                                                                                                                                                                                              'ID '
                                                                                                                                                                                              'was '
                                                                                                                                                                                              'requested:']}


def file_bytes(upload):
    """Read an upload from the start, including on Streamlit reruns."""
    if isinstance(upload, bytes):
        return upload
    if hasattr(upload, "getvalue"):
        return upload.getvalue()
    upload.seek(0)
    return upload.read()


def canonical_header(header):
    return re.sub(r"^Q\d+\s*[-–—]\s*", "", str(header).lstrip("\ufeff").strip(), flags=re.I).strip()


def read_export(upload):
    """Coalesce same-title questions, as the old export did, retaining all answers."""
    raw = file_bytes(upload)
    try:
        df = pd.read_csv(io.BytesIO(raw), dtype=str, keep_default_na=False, encoding="utf-8-sig")
    except UnicodeDecodeError:
        df = pd.read_csv(io.BytesIO(raw), dtype=str, keep_default_na=False, encoding="cp1252")
    groups = {}
    for position, header in enumerate(df.columns):
        # pandas adds .1, .2 etc. to literally duplicated CSV headers.
        base = re.sub(r"\.\d+$", "", header)
        if base in df.columns:
            header = base
        groups.setdefault(canonical_header(header), []).append(position)
    columns = {}
    mapped_titles = {canonical_header(c) for sources in QUESTION_MAPPING.values()
                     for c in (sources if isinstance(sources, list) else [sources])}
    for title, positions in groups.items():
        values = df.iloc[:, positions].replace(r"^\s*$", pd.NA, regex=True)
        conflicts = values.nunique(axis=1, dropna=True).gt(1)
        if title in mapped_titles and conflicts.any():
            raise ValueError(f"Different answers share the question title {title!r} in {int(conflicts.sum())} audit(s). Please check the export.")
        columns[title] = values.bfill(axis=1).iloc[:, 0].fillna("")
    result = pd.DataFrame(columns, index=df.index)
    required = {"internal_id", "item_to_order", "date_of_visit", "time_of_visit", "primary_result",
                "order_schedule_type", "site_code", "site_name", "site_post_code", "auditor_internal_id"}
    missing = required - set(result.columns)
    if missing:
        raise ValueError("The export is missing required columns: " + ", ".join(sorted(missing)))
    return result


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    value = str(value).strip()
    for pattern in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(value[:10] if pattern != "%d/%m/%y" else value, pattern).date()
        except ValueError:
            pass
    return None


def parse_time(value):
    if isinstance(value, time):
        return value
    for pattern in ("%H:%M:%S", "%H:%M", "%I:%M %p"):
        try:
            return datetime.strptime(str(value).strip(), pattern).time()
        except ValueError:
            pass
    return None


def most_recent_friday(today=None):
    today = today or datetime.now(ZoneInfo("Europe/London")).date()
    return today - timedelta(days=(today.weekday() - 4) % 7)


def read_previous_report(upload):
    workbook = load_workbook(io.BytesIO(file_bytes(upload)), data_only=True, keep_links=False)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"The previous report must contain the {SHEET_NAME!r} sheet.")
    sheet = workbook[SHEET_NAME]
    headers = [sheet.cell(1, c).value for c in range(1, len(REPORT_HEADERS) + 1)]
    if headers != REPORT_HEADERS:
        raise ValueError("The previous report's Weekly Data Capture columns do not match the Asda report layout.")
    ids = {str(row[0]).strip() for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True) if row[0] is not None and str(row[0]).strip()}
    if not ids and sheet.max_row > 1:
        # A header-only generated report is valid; an uncached formula report is not.
        formulas = load_workbook(io.BytesIO(file_bytes(upload)), data_only=False, read_only=True)
        has_formulas = any(c.data_type == "f" for row in formulas[SHEET_NAME].iter_rows(min_row=2, max_col=1) for c in row)
        formulas.close()
        if has_formulas:
            raise ValueError("The previous report has no saved audit IDs. Please upload the client report saved as values.")
    return workbook, ids


def select_audits(df, previous_ids, report_date):
    if report_date.weekday() != 4:
        raise ValueError("The report date must be a Friday.")
    # Monday before the previous Friday: the normal overlapping export window.
    export_start = report_date - timedelta(days=11)
    work = df.copy()
    for column in ("internal_id", "item_to_order", "primary_result", "order_schedule_type"):
        work[column] = work[column].str.strip()
    if "client_name" in work:
        client_mask = work["client_name"].str.strip().str.casefold().eq("asda stores ltd")
    else:
        # Older basic exports without client_name can still be identified by product.
        client_mask = work["item_to_order"].isin(SUPPORTED_ITEMS)
    excluded_programmes = ~client_mask | work["item_to_order"].isin({"Mystery Shop", "Shrink"})
    other_programmes = int(excluded_programmes.sum())
    work = work.loc[~excluded_programmes].copy()
    work["__date"] = work["date_of_visit"].map(parse_date)
    invalid_dates = int(work["__date"].isna().sum())
    within_window = work["__date"].map(lambda d: d is not None and export_start <= d <= report_date)
    outside_window = int((~within_window & work["__date"].notna()).sum())
    work = work.loc[within_window].copy()
    aborted = work["primary_result"].str.casefold().eq("abort")
    abort_count = int(aborted.sum())
    work = work.loc[~aborted].copy()
    already_reported = work["internal_id"].isin(previous_ids)
    previous_count = int(already_reported.sum())
    work = work.loc[~already_reported].copy()
    unsupported = sorted(set(work["item_to_order"]) - SUPPORTED_ITEMS)
    if unsupported:
        raise ValueError("These Asda products need a report mapping before they can be included: " + ", ".join(unsupported))
    if work["internal_id"].eq("").any():
        raise ValueError("An eligible audit is missing its internal_id, so it cannot safely be checked against last week's report.")
    invalid_results = ~work["primary_result"].str.casefold().isin({"pass", "fail"})
    if invalid_results.any():
        raise ValueError("Eligible audits have results other than pass/fail. Please check: " + ", ".join(work.loc[invalid_results, "internal_id"].head(10)))
    duplicates = work.duplicated("internal_id", keep=False)
    if duplicates.any():
        for audit_id, group in work.loc[duplicates].groupby("internal_id"):
            if len(group.drop_duplicates()) > 1:
                raise ValueError(f"Audit {audit_id} appears more than once with different data. Please check the export.")
    duplicate_count = int(work.duplicated("internal_id").sum())
    work = work.drop_duplicates("internal_id").copy()
    work["__time"] = work["time_of_visit"].map(parse_time)
    if work["__time"].isna().any():
        raise ValueError("An eligible audit has an invalid time_of_visit. Please check: " + ", ".join(work.loc[work["__time"].isna(), "internal_id"].head(10)))
    if (~work["order_schedule_type"].isin(VISIT_TYPES)).any():
        unknown = sorted(set(work["order_schedule_type"]) - set(VISIT_TYPES))
        raise ValueError("These schedule types need a Visit Type mapping: " + ", ".join(unknown))
    work = work.sort_values(["item_to_order", "__date", "__time"], kind="stable").reset_index(drop=True)
    stats = {"Export rows": len(df), "Other programmes": other_programmes,
             "Outside reporting window": outside_window, "Invalid visit dates": invalid_dates,
             "Aborts": abort_count, "Already in previous report": previous_count,
             "Repeated export rows": duplicate_count, "New audits": len(work)}
    return work, stats, export_start


def text(value):
    return "" if value is None or pd.isna(value) else str(value)


def first_answer(row, source):
    sources = source if isinstance(source, list) else [source]
    return next((text(row.get(c, "")) for c in sources if text(row.get(c, "")).strip()), "")


def excel_value(value):
    """Keep alphanumeric/zero-prefixed receipt codes as text, numeric answers as numbers."""
    value = text(value)
    if re.fullmatch(r"-?(0|[1-9]\d*)(\.\d+)?", value.strip()):
        number = Decimal(value.strip())
        return int(number) if number == number.to_integral_value() else float(number)
    return value


def proper_store_name(value):
    # Excel PROPER capitalises after every non-letter, including digits/apostrophes.
    value = re.sub(r"[A-Za-z]+", lambda m: m.group().capitalize(), text(value))
    for old, new in (("Pfs", "PFS"), ("Iow", "IOW"), ("Asda - ", ""),
                     ("Asda, ", ""), ("Asda Express - ", ""), ("Asda Express, ", ""),
                     ("Asda Express ", ""), ("Asda ", "")):
        value = value.replace(old, new)
    return value


def report_rows(audits):
    rows = []
    for _, source in audits.iterrows():
        row = source.to_dict()
        item = row["item_to_order"]
        if CHECKOUT_ID in audits.columns or CABINET_ID in audits.columns:
            row["merged_id_asked"] = "Yes" if any(text(row.get(c, "")) == "Yes" for c in (CHECKOUT_ID, CABINET_ID)) else "No"
        if item == "Fireworks - No ID":
            allow, handover = text(row.get(ALLOW, "")), text(row.get(HANDOVER, ""))
            if ALLOW in audits.columns and HANDOVER in audits.columns:
                row["merged_no_id_allow_handover"] = allow if allow == "No" else handover
            else:
                row["merged_no_id_allow_handover"] = allow or handover
            row[ALLOW] = ""
        answers = {question: first_answer(row, sources) for question, sources in QUESTION_MAPPING.items()}
        def answer(question):
            return answers.get(question, "")
        asked = answer("Did the store colleague who served you ask you for ID?")
        bought = answer(ALLOW)
        if not bought and item in DELIVERY_ITEMS:
            bought = "No" if asked == "Yes" else "Yes"
        till = answer("At which type of till was the purchase made?")
        if not till and item not in DELIVERY_ITEMS | {"Fireworks - No ID"}:
            till = "Kiosk"
        description = answer("Please input the approximate height of the server:")
        for question, prefix, suffix in (
            ("Please input the length of the server's hair:", ", ", ""),
            ("Please input the style of the server hair:", " ", ""),
            ("Please input the hair colour of the server:", " ", " hair"),
            ("Please select whether the server had any of the following features:", ", ", ""),
            ("Please input their distinguishing characteristics:", ", ", ""),
        ):
            value = answer(question)
            if value and not (question.startswith("Please select") and value == "None of the above"):
                description += prefix + value + suffix
        queue = excel_value(answer("How many people were in the queue?"))
        code = text(row.get("site_code", ""))[:4]
        try:
            location = int(code)
        except ValueError:
            raise ValueError(f"Audit {row['internal_id']} has an invalid site_code: {row.get('site_code')!r}.")
        values = [
            row["internal_id"], VISIT_TYPES[row["order_schedule_type"]], item,
            datetime.combine(row["__date"], time()), row["__time"], location,
            proper_store_name(row["site_name"]), row["site_post_code"], asked, bought, till,
            ("No" if queue == 0 else "Yes") if queue != "" else "", queue if queue != 0 else "",
            answer("Did the store colleague who served you make eye contact with you during the transaction?") + " " + answer("When was eye contact first made?"),
            "", answer("What was the gender of the store colleague who served you?"),
            answer("What was the approximate age of the store colleague who served you?"),
            answer("What was the name of the store colleague who served you?"), description,
            answer("Please give details of the product that you purchased:"),
            excel_value(answer("From the receipt, please enter the 4-digit ST code:")),
            excel_value(answer("Please enter the 8-digit OP code:")),
            excel_value(answer("Please enter the 2-digit TE code:")),
            excel_value(answer("Please enter the 5-digit TR code:")),
            answer("Please use this space to explain anything unusual about your visit or to clarify any detail of your report:"),
            row["auditor_internal_id"], answer("Did the staff member approach you to attend to your till?"),
        ]
        rows.append(values)
    return rows


def write_report(workbook, rows):
    """Reuse the previous client report's layout; replace all data with computed values."""
    sheet = workbook[SHEET_NAME]
    styles = [copy(sheet.cell(2, c)._style) for c in range(1, len(REPORT_HEADERS) + 1)]
    row_height = sheet.row_dimensions[2].height
    for name in list(workbook.sheetnames):
        if name != SHEET_NAME:
            del workbook[name]
    # No formulas or named ranges referring to removed LIVE sheets are retained.
    workbook.defined_names.clear()
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for name in list(sheet.tables):
        del sheet.tables[name]
    for row_num, values in enumerate(rows, start=2):
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_num, column)
            cell._style = copy(styles[column - 1])
            cell.value = value if value != "" else None
            # User-provided text beginning '=' remains text, never an Excel formula.
            if isinstance(value, str) and value:
                cell.data_type = "s"
            if column == 4:
                cell.number_format = "mm-dd-yy"
            elif column == 5:
                cell.number_format = "h:mm"
        if row_height is not None:
            sheet.row_dimensions[row_num].height = row_height
    sheet.auto_filter.ref = f"A1:AA{len(rows) + 1}"
    if sheet.print_area:
        sheet.print_area = f"A1:AA{len(rows) + 1}"
    workbook.active = 0
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def percentage(passes, total):
    if not total:
        return "0%"
    return str((Decimal(passes) * 100 / Decimal(total)).quantize(Decimal("1"), rounding=ROUND_HALF_UP)) + "%"


def pass_rate_summary(audits):
    summary = []
    for item, group in audits.groupby("item_to_order", sort=True):
        count = len(group)
        passes = int(group["primary_result"].str.casefold().eq("pass").sum())
        summary.append({"Product": item, "Audits": count, "Passes": passes, "Pass rate": percentage(passes, count)})
    return summary


def email_text(audits, report_date):
    day = report_date.day
    suffix = "ᵗʰ" if 11 <= day <= 13 else {1: "ˢᵗ", 2: "ⁿᵈ", 3: "ʳᵈ"}.get(day % 10, "ᵗʰ")
    count = len(audits)
    passes = int(audits["primary_result"].str.casefold().eq("pass").sum())
    message = ("Hi all,\n\n"
               f"Please find attached the Serve Legal weekly report detailing visits carried out up to Friday, {day}{suffix} {report_date:%B}.\n\n"
               f"As you’ll see from the attached, the pass rate was {percentage(passes, count)} based on {count} completed audit{'s' if count != 1 else ''}. This is broken down as below:\n")
    for summary in pass_rate_summary(audits):
        count = summary["Audits"]
        message += f"\n{summary['Product']} – {count} visit{'s' if count != 1 else ''} ({summary['Pass rate']})"
    return message


def generate_report(export_upload, previous_upload, report_date):
    """Pure entry point for testing and reuse outside Streamlit."""
    export = read_export(export_upload)
    workbook, previous_ids = read_previous_report(previous_upload)
    audits, stats, start = select_audits(export, previous_ids, report_date)
    rows = report_rows(audits)
    return {"report": write_report(workbook, rows), "email": email_text(audits, report_date),
            "rows": rows, "audits": audits, "stats": stats, "start": start,
            "summary": pass_rate_summary(audits),
            "filename": f"Asda Weekly {report_date:%d.%m.%Y}.xlsx"}


def main():
    import streamlit as st
    st.title("Asda Weekly Report Generator")
    st.write("Upload an audit export and the previous week's client report to create the updated report and email text.")
    export = st.file_uploader("Audit export (old or new format)", type=["csv"])
    previous = st.file_uploader("Previous week's report", type=["xlsx"])
    report_date = st.date_input("Report date (Friday)", value=most_recent_friday())
    st.caption("Visits are checked from the Monday before the previous report through the selected Friday. Audits already in the previous report, aborts, Mystery Shop and Shrink are excluded.")
    if report_date.weekday() != 4:
        st.error("Please select a Friday as the report date.")
        return
    st.caption(f"Export window: {(report_date - timedelta(days=11)):%d/%m/%Y} to {report_date:%d/%m/%Y}, inclusive.")
    if export is None or previous is None:
        return
    try:
        result = generate_report(export, previous, report_date)
    except Exception as error:
        st.error(f"The report could not be generated: {error}")
        return
    stats = result["stats"]
    if stats["Invalid visit dates"]:
        st.warning(f"{stats['Invalid visit dates']} Asda audit(s) had an invalid/missing visit date and were excluded. Check the export before sending the report.")
    if not stats["New audits"]:
        st.info("No new completed audits were found for this reporting window.")
    else:
        st.success(f"Report ready: {stats['New audits']} new completed audits.")
    with st.expander("Audit checks"):
        st.dataframe(pd.DataFrame(list(stats.items()), columns=["Check", "Rows"]), hide_index=True, use_container_width=True)
    if result["summary"]:
        total = len(result["audits"])
        passes = int(result["audits"]["primary_result"].str.casefold().eq("pass").sum())
        st.metric("Overall pass rate", percentage(passes, total))
        st.dataframe(pd.DataFrame(result["summary"])[["Product", "Audits", "Pass rate"]], hide_index=True, use_container_width=True)
    st.download_button("Download updated report", result["report"], result["filename"],
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    st.text_area("Email text", result["email"], height=360)
    st.download_button("Download email text", result["email"].encode("utf-8-sig"),
                       f"Asda Weekly {report_date:%d.%m.%Y} Email.txt", "text/plain")
    with st.expander("Report preview"):
        st.dataframe(pd.DataFrame(result["rows"], columns=REPORT_HEADERS), hide_index=True, use_container_width=True)


if __name__ == "__main__":
    main()
